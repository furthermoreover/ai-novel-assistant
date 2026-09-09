# -*- coding: utf-8 -*-
"""小说业务服务：项目管理、大纲、人物、章节、AI 写作"""
import io
import os
import re
import uuid
from db import get_conn
from ai import chat_completion
from knowledge import search_knowledge
from skills import WRITING_SKILLS, SKILL_MAP, match_skill

# ---------- 系统提示词（从 prompts/system_prompt.md 读取，可自行编辑） ----------
_SYSTEM_PROMPT_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "prompts", "system_prompt.md",
)
_SYSTEM_PROMPT_CACHE = None


def get_system_prompt() -> str:
    """读取小说助手系统提示词（带缓存，编辑 md 后重启生效）"""
    global _SYSTEM_PROMPT_CACHE
    if _SYSTEM_PROMPT_CACHE is None:
        try:
            with open(_SYSTEM_PROMPT_PATH, "r", encoding="utf-8") as f:
                _SYSTEM_PROMPT_CACHE = f.read().strip()
        except Exception:
            _SYSTEM_PROMPT_CACHE = (
                "你是一位专业的网络小说创作助手，擅长各类题材，"
                "严格遵循大纲与人设，输出连贯的正文。"
            )
    return _SYSTEM_PROMPT_CACHE

XUANHUAN_KEYWORDS = [
    "玄幻", "仙侠", "修真", "修仙", "修炼", "境界", "金丹", "元婴",
    "宗门", "法宝", "妖兽", "灵根", "废柴", "逆袭", "系统", "签到",
    "洪荒", "渡劫", "飞升", "炼气", "筑基", "化神", "洞天", "符箓",
    "丹药", "天材地宝", "灵宠", "全球神祇",
]


def _count_words(content: str) -> int:
    if not content:
        return 0
    cn = len(re_findall(r"[\u4e00-\u9fa5]", content))
    en = len(re_findall(r"[a-zA-Z]+", content))
    return cn + en


def re_findall(pattern, text):
    return re.findall(pattern, text)


# ---------- 项目 ----------
def list_projects(user_id: str) -> dict:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM novel_project WHERE user_id=%s ORDER BY _created_at DESC", (user_id,))
            items = cur.fetchall()
            cur.execute("SELECT COUNT(*) AS total FROM novel_project WHERE user_id=%s", (user_id,))
            total = cur.fetchone()["total"]
    return {"items": items, "total": total}


def create_project(title: str, genre: str = "", description: str = "", user_id: str | None = None) -> dict:
    pid = str(uuid.uuid4())
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO novel_project (id, title, genre, description, user_id) VALUES (%s,%s,%s,%s,%s) RETURNING *",
                (pid, title, genre, description, user_id),
            )
            project = cur.fetchone()
            # 默认大纲
            cur.execute(
                "INSERT INTO novel_outline (project_id) VALUES (%s)", (pid,),
            )
        conn.commit()
    return project


def get_project(pid: str) -> dict | None:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM novel_project WHERE id=%s", (pid,))
            return cur.fetchone()


def update_project(pid: str, **fields) -> dict:
    allowed = {"title", "genre", "description", "cover_image"}
    sets = []
    params = []
    for k, v in fields.items():
        if k in allowed and v is not None:
            sets.append(f"{k}=%s")
            params.append(v)
    if not sets:
        return get_project(pid)
    params.append(pid)
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(f"UPDATE novel_project SET {', '.join(sets)}, _updated_at=CURRENT_TIMESTAMP WHERE id=%s RETURNING *", params)
            row = cur.fetchone()
        conn.commit()
    return row


def delete_project(pid: str) -> None:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM novel_project WHERE id=%s", (pid,))
        conn.commit()


# ---------- 大纲 ----------
def get_outline(pid: str) -> dict | None:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM novel_outline WHERE project_id=%s", (pid,))
            return cur.fetchone()


def upsert_outline(pid: str, **fields) -> dict:
    allowed = ["main_storyline", "world_setting", "volume_plans", "key_plot_points", "foreshadowing", "other_notes"]
    existing = get_outline(pid)
    sets = ["_updated_at=CURRENT_TIMESTAMP"]
    params = []
    for k in allowed:
        if k in fields and fields[k] is not None:
            sets.append(f"{k}=%s")
            params.append(fields[k])
    params.append(pid)
    with get_conn() as conn:
        with conn.cursor() as cur:
            if existing:
                cur.execute(f"UPDATE novel_outline SET {', '.join(sets)} WHERE project_id=%s RETURNING *", params)
            else:
                cols = ["project_id"] + [k for k in allowed if k in fields and fields[k] is not None]
                vals = [pid] + [fields[k] for k in allowed if k in fields and fields[k] is not None]
                cur.execute(
                    f"INSERT INTO novel_outline ({', '.join(cols)}) VALUES ({', '.join(['%s'] * len(vals))}) RETURNING *",
                    vals,
                )
            row = cur.fetchone()
        conn.commit()
    return row


# ---------- 人物 ----------
def list_characters(pid: str) -> list[dict]:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT * FROM novel_character WHERE project_id=%s ORDER BY sort_order, _created_at",
                (pid,),
            )
            return cur.fetchall()


def create_character(pid: str, name: str, **fields) -> dict:
    cid = str(uuid.uuid4())
    cols = ["id", "project_id", "name"]
    vals = [cid, pid, name]
    allowed = ["age", "identity", "appearance", "personality", "background", "motivation",
               "relationships", "speech_style", "avatar", "other_info", "sort_order"]
    for k in allowed:
        if k in fields and fields[k] is not None:
            cols.append(k)
            vals.append(fields[k])
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"INSERT INTO novel_character ({', '.join(cols)}) VALUES ({', '.join(['%s'] * len(vals))}) RETURNING *",
                vals,
            )
            row = cur.fetchone()
        conn.commit()
    return row


def update_character(pid: str, cid: str, **fields) -> dict | None:
    allowed = ["name", "age", "identity", "appearance", "personality", "background", "motivation",
               "relationships", "speech_style", "avatar", "other_info", "sort_order"]
    sets = []
    params = []
    for k in allowed:
        if k in fields and fields[k] is not None:
            sets.append(f"{k}=%s")
            params.append(fields[k])
    if not sets:
        return None
    sets.append("_updated_at=CURRENT_TIMESTAMP")
    params += [pid, cid]
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"UPDATE novel_character SET {', '.join(sets)} WHERE project_id=%s AND id=%s RETURNING *",
                params,
            )
            row = cur.fetchone()
        conn.commit()
    return row


def delete_character(pid: str, cid: str) -> None:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM novel_character WHERE project_id=%s AND id=%s", (pid, cid))
        conn.commit()


# ---------- 章节 ----------
def list_chapters(pid: str) -> dict:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM novel_chapter WHERE project_id=%s ORDER BY chapter_number", (pid,))
            items = cur.fetchall()
            cur.execute("SELECT COUNT(*) AS total FROM novel_chapter WHERE project_id=%s", (pid,))
            total = cur.fetchone()["total"]
    return {"items": items, "total": total}


def create_chapter(pid: str, title: str, content: str = "", chapter_number: int | None = None) -> dict:
    with get_conn() as conn:
        with conn.cursor() as cur:
            if chapter_number and chapter_number > 0:
                num = chapter_number
                # 防重复章号：若该章号已存在（如重复点击"按规划新建章节"），自动顺延到下一个可用编号
                cur.execute("SELECT 1 FROM novel_chapter WHERE project_id=%s AND chapter_number=%s", (pid, num))
                if cur.fetchone():
                    cur.execute("SELECT COALESCE(MAX(chapter_number), 0) + 1 AS n FROM novel_chapter WHERE project_id=%s", (pid,))
                    num = cur.fetchone()["n"]
            else:
                cur.execute("SELECT COALESCE(MAX(chapter_number), 0) + 1 AS n FROM novel_chapter WHERE project_id=%s", (pid,))
                num = cur.fetchone()["n"]
            cid = str(uuid.uuid4())
            cur.execute(
                """INSERT INTO novel_chapter (id, project_id, chapter_number, title, content, word_count)
                   VALUES (%s,%s,%s,%s,%s,%s) RETURNING *""",
                (cid, pid, num, title, content, _count_words(content)),
            )
            row = cur.fetchone()
        conn.commit()
    return row


def get_chapter(pid: str, cid: str) -> dict | None:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM novel_chapter WHERE project_id=%s AND id=%s", (pid, cid))
            return cur.fetchone()


def update_chapter(pid: str, cid: str, **fields) -> dict | None:
    allowed = ["title", "content", "summary", "status"]
    sets = []
    params = []
    for k in allowed:
        if k in fields and fields[k] is not None:
            sets.append(f"{k}=%s")
            params.append(fields[k])
            if k == "content":
                sets.append("word_count=%s")
                params.append(_count_words(fields[k]))
    if not sets:
        return None
    sets.append("_updated_at=CURRENT_TIMESTAMP")
    params += [pid, cid]
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"UPDATE novel_chapter SET {', '.join(sets)} WHERE project_id=%s AND id=%s RETURNING *",
                params,
            )
            row = cur.fetchone()
        conn.commit()
    return row


def _purge_chapter_memory(pid: str, num: int) -> None:
    """删除章节时，级联清理该章在 Agent 记忆库中的残留，防止废稿被记忆：
    - 剧情时间线：该章记录整条删除
    - 伏笔：埋设于该章（来源消失）→ 删除；回收于该章（回收事件不存在）→ 恢复为未回收
    - 角色卡：最近出场章节指向该章 → 置空（角色卡本身保留）
    """
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM novel_timeline WHERE project_id=%s AND chapter_number=%s", (pid, num))
            cur.execute("DELETE FROM novel_foreshadowing WHERE project_id=%s AND planted_chapter=%s", (pid, num))
            cur.execute(
                "UPDATE novel_foreshadowing SET status='open', resolved_chapter=NULL WHERE project_id=%s AND resolved_chapter=%s",
                (pid, num),
            )
            cur.execute(
                "UPDATE novel_character SET last_appearance_chapter=NULL WHERE project_id=%s AND last_appearance_chapter=%s",
                (pid, num),
            )
        conn.commit()


def delete_chapter(pid: str, cid: str) -> None:
    num = None
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT chapter_number FROM novel_chapter WHERE project_id=%s AND id=%s", (pid, cid))
            row = cur.fetchone()
            num = row["chapter_number"] if row else None
            cur.execute("DELETE FROM novel_chapter WHERE project_id=%s AND id=%s", (pid, cid))
        conn.commit()
    if num is not None:
        _purge_chapter_memory(pid, num)


def cleanup_orphan_memory(pid: str) -> dict:
    """清理孤儿记忆：删除/修复指向"已不存在章节"的记忆记录。

    场景：历史版本删除章节时未清理记忆库，导致废稿（如被删的第 29 章）
    仍残留在时间线/伏笔中。幂等，可安全重复执行。
    """
    counts = {"timeline": 0, "foreshadowing_deleted": 0, "foreshadowing_reopened": 0, "characters_fixed": 0}
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT chapter_number FROM novel_chapter WHERE project_id=%s", (pid,))
            nums = [r["chapter_number"] for r in cur.fetchall()]
            if nums:
                cur.execute(
                    "DELETE FROM novel_timeline WHERE project_id=%s AND NOT (chapter_number = ANY(%s))",
                    (pid, nums),
                )
                counts["timeline"] = cur.rowcount
                cur.execute(
                    "DELETE FROM novel_foreshadowing WHERE project_id=%s AND planted_chapter IS NOT NULL AND NOT (planted_chapter = ANY(%s))",
                    (pid, nums),
                )
                counts["foreshadowing_deleted"] = cur.rowcount
                cur.execute(
                    "UPDATE novel_foreshadowing SET status='open', resolved_chapter=NULL WHERE project_id=%s AND resolved_chapter IS NOT NULL AND NOT (resolved_chapter = ANY(%s))",
                    (pid, nums),
                )
                counts["foreshadowing_reopened"] = cur.rowcount
                cur.execute(
                    "UPDATE novel_character SET last_appearance_chapter=NULL WHERE project_id=%s AND last_appearance_chapter IS NOT NULL AND NOT (last_appearance_chapter = ANY(%s))",
                    (pid, nums),
                )
                counts["characters_fixed"] = cur.rowcount
            else:
                # 项目没有任何章节：清空时间线（伏笔无从归属，一并清空）
                cur.execute("DELETE FROM novel_timeline WHERE project_id=%s", (pid,))
                counts["timeline"] = cur.rowcount
                cur.execute("DELETE FROM novel_foreshadowing WHERE project_id=%s", (pid,))
                counts["foreshadowing_deleted"] = cur.rowcount
        conn.commit()
    return counts


def cleanup_all_orphan_memory() -> dict:
    """对所有项目执行孤儿记忆清理（服务启动时调用，幂等）"""
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM novel_project")
            pids = [r["id"] for r in cur.fetchall()]
    total = {"projects": len(pids), "timeline": 0, "foreshadowing_deleted": 0,
             "foreshadowing_reopened": 0, "characters_fixed": 0}
    for pid in pids:
        c = cleanup_orphan_memory(pid)
        for k in ("timeline", "foreshadowing_deleted", "foreshadowing_reopened", "characters_fixed"):
            total[k] += c[k]
    return total


# 章节标题正则：匹配「第X章/回/节/卷 + 标题」
CHAPTER_TITLE_RE = re.compile(r"^\s*(第[一二三四五六七八九十百千万零〇0-9０-９]+[章回节卷][^\n]{0,50})\s*$", re.MULTILINE)


def import_chapters(pid: str, content: str) -> dict:
    """把已有小说文本按章节标题自动拆分成章节导入"""
    content = content.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not content:
        return {"imported": 0, "chapters": [], "message": "内容为空"}

    # 找章节标题位置，按标题切分
    matches = list(CHAPTER_TITLE_RE.finditer(content))
    chunks: list[tuple[str, str]] = []
    if not matches:
        # 没有章节标题，整体作为一章
        chunks = [("第一章", content)]
    else:
        for i, m in enumerate(matches):
            title = m.group(1).strip()
            start = m.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(content)
            body = content[start:end].strip()
            chunks.append((title, body))

    # 已有章节数决定起始编号
    existing = list_chapters(pid)["total"]

    imported = []
    with get_conn() as conn:
        with conn.cursor() as cur:
            for idx, (title, body) in enumerate(chunks):
                num = existing + idx + 1
                cid = str(uuid.uuid4())
                cur.execute(
                    """INSERT INTO novel_chapter
                       (id, project_id, chapter_number, title, content, word_count, status)
                       VALUES (%s,%s,%s,%s,%s,%s,'published') RETURNING *""",
                    (cid, pid, num, title[:255], body, _count_words(body)),
                )
                imported.append(cur.fetchone())
        conn.commit()
    return {"imported": len(imported), "chapters": imported, "message": f"成功导入 {len(imported)} 章"}


# ---------- AI 生成 ----------
def _build_context(pid: str, current_chapter_id: str | None = None):
    outline = get_outline(pid)
    chars = list_characters(pid)
    with get_conn() as conn:
        with conn.cursor() as cur:
            if current_chapter_id:
                cur.execute("SELECT chapter_number FROM novel_chapter WHERE id=%s", (current_chapter_id,))
                row = cur.fetchone()
                cur_num = row["chapter_number"] if row else 999999
                cur.execute(
                    "SELECT * FROM novel_chapter WHERE project_id=%s AND chapter_number < %s ORDER BY chapter_number DESC LIMIT 3",
                    (pid, cur_num),
                )
            else:
                cur.execute("SELECT * FROM novel_chapter WHERE project_id=%s ORDER BY chapter_number DESC LIMIT 3", (pid,))
            recent = list(reversed(cur.fetchall()))
    return outline, chars, recent


def _outline_text(outline) -> str:
    if not outline:
        return ""
    parts = []
    labels = [("main_storyline", "主线剧情"), ("world_setting", "世界设定"), ("volume_plans", "分卷规划"),
              ("key_plot_points", "关键情节"), ("foreshadowing", "伏笔"), ("other_notes", "其他备注")]
    for key, label in labels:
        if outline.get(key):
            parts.append(f"【{label}】\n{outline[key]}")
    return "\n\n".join(parts)


def _characters_text(chars) -> str:
    if not chars:
        return ""
    blocks = []
    for c in chars:
        lines = [f"===== {c['name']} ====="]
        for key, label in [("age", "年龄"), ("identity", "身份"), ("personality", "性格"),
                           ("background", "背景"), ("motivation", "动机"), ("relationships", "人物关系"),
                           ("speech_style", "说话风格"), ("appearance", "外貌")]:
            if c.get(key):
                lines.append(f"{label}：{c[key]}")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def _recent_text(recent) -> str:
    if not recent:
        return ""
    return "\n\n".join(
        f"第{c['chapter_number']}章 {c['title']}\n摘要：{c.get('summary') or ''}" for c in recent
    )


def generate_chapter(pid: str, cid: str, instruction: str,
                     skill_id: str | None = None, skill_mode: str = "standard",
                     selected_topic: str | None = None,
                     state_context: str | None = None):
    """AI 生成章节正文（技能 + 知识库协同，三档模式）。

    state_context: 若传入 Agent 作品状态包（精简/完整），则用它替代
    大纲全文+人物卡全文+最近3章摘要，控制 token 且记忆更全。
    """
    project = get_project(pid)
    chapter = get_chapter(pid, cid)
    if not chapter:
        raise ValueError("章节不存在")
    genre = (project or {}).get("genre") or ""
    outline, chars, recent = _build_context(pid, cid)

    # 匹配技能
    skill = None
    if skill_id and skill_id != "auto":
        skill = SKILL_MAP.get(skill_id)
    if not skill:
        skill = match_skill(instruction)

    # 技能 + 知识库协同
    chunks = []
    if skill_mode != "skill-only":
        k_limit = 4 if skill_mode == "deep" else 2
        if skill and skill["boundTopics"]:
            topics = list(skill["boundTopics"])
            # 玄幻识别
            text = instruction + " " + genre
            if any(kw in text for kw in XUANHUAN_KEYWORDS) and "玄幻题材专项" not in topics:
                topics.append("玄幻题材专项")
            per_topic = max(1, -(-k_limit // len(topics)))
            seen = set()
            for topic in topics:
                if len(chunks) >= k_limit:
                    break
                for c in search_knowledge(instruction, per_topic, topic):
                    if c["id"] not in seen and len(chunks) < k_limit:
                        seen.add(c["id"])
                        chunks.append(c)
        elif selected_topic and selected_topic != "auto":
            chunks = search_knowledge(instruction, k_limit, selected_topic)
        else:
            chunks = search_knowledge(instruction, k_limit)

    # 组装提示词
    if skill:
        knowledge_text = "\n\n".join(
            f"{i+1}. {c['title']}（{c['topic']}）\n{c['content'][:300]}" for i, c in enumerate(chunks)
        ) if chunks else "（无知识库参考，仅按技能模板生成）"
        if state_context:
            # Agent 状态包模式：注入精简/完整状态包，不再重复注入大纲全文
            ot = state_context
            ct = "（已并入上方作品状态包）"
            rt = "（已并入上方作品状态包）"
        else:
            ot = _outline_text(outline) or "（暂无大纲）"
            ct = _characters_text(chars) or "（暂无人设）"
            rt = _recent_text(recent) or "（无已写章节）"
        prompt = (skill["promptTemplate"]
                  .replace("{{instruction}}", instruction)
                  .replace("{{knowledge}}", knowledge_text)
                  .replace("{{outline}}", ot)
                  .replace("{{characters}}", ct)
                  .replace("{{recent_summary}}", rt))
    else:
        prompt = instruction
        if chunks:
            ref = "\n\n".join(f"{i+1}. {c['title']}（主题：{c['topic']}）\n{c['content'][:200]}" for i, c in enumerate(chunks))
            prompt = f"【写作技法参考】\n{ref}\n\n【创作指令】\n{instruction}"

    messages = [
        {"role": "system", "content": get_system_prompt()},
        {"role": "user", "content": prompt},
    ]

    generated = chat_completion(messages, temperature=0.8, max_tokens=4096)

    # 追加保存
    new_content = (chapter.get("content") or "") + generated
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE novel_chapter SET content=%s, word_count=%s, _updated_at=CURRENT_TIMESTAMP WHERE id=%s RETURNING *",
                (new_content, _count_words(new_content), cid),
            )
            row = cur.fetchone()
        conn.commit()
    return row


def summarize_chapter(pid: str, cid: str) -> dict:
    """AI 生成章节摘要"""
    chapter = get_chapter(pid, cid)
    if not chapter:
        raise ValueError("章节不存在")
    content = chapter.get("content") or ""
    if not content:
        raise ValueError("章节内容为空，无法生成摘要")
    messages = [
        {"role": "system", "content": "你是小说章节摘要助手，输出 150 字以内的精炼摘要，包含本章关键事件、人物动向、埋下的伏笔。"},
        {"role": "user", "content": f"请为以下章节生成摘要：\n\n{content[:8000]}"},
    ]
    summary = chat_completion(messages, temperature=0.3, max_tokens=500)
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE novel_chapter SET summary=%s, _updated_at=CURRENT_TIMESTAMP WHERE id=%s RETURNING *",
                (summary.strip(), cid),
            )
            row = cur.fetchone()
        conn.commit()
    return row


def list_skills() -> list[dict]:
    return WRITING_SKILLS


# ---------- 导出 ----------
def export_project(pid: str, fmt: str = "md", chapter_from: int | None = None, chapter_to: int | None = None) -> dict:
    """批量导出小说：fmt 支持 md / txt / docx；chapter_from / chapter_to 限定范围（全部导出时省略）。

    返回 {"filename": str, "text": str}（md/txt）或 {"filename": str, "bytes": bytes}（docx）。
    """
    chapters = list_chapters(pid)["items"]
    selected = []
    for c in chapters:
        n = c["chapter_number"]
        if chapter_from is not None and n < chapter_from:
            continue
        if chapter_to is not None and n > chapter_to:
            continue
        selected.append(c)
    if not selected:
        raise ValueError("所选范围内没有可导出的章节")

    project = get_project(pid)
    title = ((project or {}).get("title") or "小说").strip() or "小说"
    fmt = (fmt or "md").lower()

    if fmt == "docx":
        try:
            from docx import Document
        except ImportError:
            raise ValueError("导出 docx 需要安装 python-docx（pip install python-docx）")
        doc = Document()
        doc.add_heading(title, 0)
        for c in selected:
            doc.add_heading(_export_chapter_head(c), level=1)
            doc.add_paragraph(c["content"] or "")
        buf = io.BytesIO()
        doc.save(buf)
        return {"filename": f"{title}.docx", "bytes": buf.getvalue()}

    # txt / md：按章节顺序拼接
    parts = []
    if fmt == "md":
        parts.append(f"# {title}")
    else:
        parts.append(title)
        parts.append("=" * len(title))
    for c in selected:
        head = _export_chapter_head(c)
        if fmt == "md":
            parts.append(f"## {head}")
        else:
            parts.append(head)
        parts.append(c["content"] or "")
    text = "\n\n".join(parts)
    return {"filename": f"{title}.{fmt}", "text": text}


def _export_chapter_head(c: dict) -> str:
    """章节标题：若标题已带「第X章/回」前缀则原样使用，否则自动补章号，避免导出重复"""
    t = (c.get("title") or "").strip()
    if re.match(r"^第[0-9一二三四五六七八九十百千万零〇]+[章回节卷]", t):
        return t
    return f"第{c['chapter_number']}章 {t}"
